import os
import io
import streamlit as st
from dotenv import load_dotenv

# Document libraries
import fitz  # PyMuPDF
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.pdfgen import canvas
from reportlab.lib.units import inch

# AI and Image libraries
import PIL.Image
import google.generativeai as genai

# Load env variables
load_dotenv()

# Set page config
st.set_page_config(
    page_title="AI Event Report Generator",
    page_icon="📋",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ---------------------------------------------------------
# DOCUMENT EXTRACTION & EXPORT MODULES (Inline report_exporter.py)
# ---------------------------------------------------------

def extract_text_from_pdf(file_bytes: bytes) -> str:
    """Extracts text from PDF file bytes."""
    text_content = []
    try:
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        for page in doc:
            text = page.get_text()
            if text:
                text_content.append(text)
    except Exception as e:
        return f"Error reading PDF: {str(e)}"
    return "\n".join(text_content)

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extracts text from DOCX file bytes."""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        text_content = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_content.append(para.text)
        return "\n".join(text_content)
    except Exception as e:
        return f"Error reading DOCX: {str(e)}"


class NumberedCanvas(canvas.Canvas):
    """ReportLab Canvas that automatically adds dynamic header, footer, and page counts."""
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_elements(num_pages)
            super().showPage()
        super().save()

    def draw_page_elements(self, page_count):
        # Do not draw headers/footers on page 1 (Cover Page)
        if self._pageNumber == 1:
            return
        
        self.saveState()
        self.setFont("Helvetica-Bold", 8)
        self.setFillColor(colors.HexColor("#475569"))
        
        # Header text
        self.drawString(54, 755, "EVENT COMPREHENSIVE REPORT")
        self.setFont("Helvetica", 8)
        self.drawRightString(558, 755, "AI Event Report Generator")
        
        # Header line
        self.setStrokeColor(colors.HexColor("#cbd5e1"))
        self.setLineWidth(0.75)
        self.line(54, 747, 558, 747)
        
        # Footer line
        self.line(54, 55, 558, 55)
        
        # Footer text
        self.drawString(54, 42, "Generated automatically by AI Report Engine")
        page_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(558, 42, page_text)
        
        self.restoreState()


def generate_pdf(report_data: dict, logo_bytes: bytes = None, image_categories: dict = None) -> bytes:
    """Generates a professional PDF report."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=72,
        bottomMargin=72
    )
    
    styles = getSampleStyleSheet()
    
    # Custom Styles
    primary_color = colors.HexColor("#1e3a8a")  # Deep blue
    secondary_color = colors.HexColor("#0d9488")  # Teal
    text_color = colors.HexColor("#1e293b")  # Dark slate
    light_bg = colors.HexColor("#f8fafc")
    
    title_style = ParagraphStyle(
        'CoverTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=26,
        leading=32,
        textColor=primary_color,
        alignment=1, # Centered
        spaceAfter=15
    )
    
    subtitle_style = ParagraphStyle(
        'CoverSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=14,
        leading=18,
        textColor=secondary_color,
        alignment=1,
        spaceAfter=30
    )
    
    h1_style = ParagraphStyle(
        'Heading1_Custom',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        textColor=primary_color,
        spaceBefore=16,
        spaceAfter=8,
        keepWithNext=True
    )
    
    h2_style = ParagraphStyle(
        'Heading2_Custom',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        textColor=secondary_color,
        spaceBefore=10,
        spaceAfter=6,
        keepWithNext=True
    )
    
    body_style = ParagraphStyle(
        'Body_Custom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=text_color,
        spaceAfter=8
    )
    
    table_label_style = ParagraphStyle(
        'TableLabel',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=12,
        textColor=primary_color
    )
    
    story = []
    
    # ---------------------------------------------------------
    # COVER PAGE
    # ---------------------------------------------------------
    story.append(Spacer(1, 40))
    
    # Logo placement if available
    if logo_bytes:
        try:
            logo_img = Image(io.BytesIO(logo_bytes), width=1.5*inch, height=1.5*inch)
            logo_img.hAlign = 'CENTER'
            story.append(logo_img)
            story.append(Spacer(1, 25))
        except Exception:
            pass
            
    story.append(Paragraph("EVENT COMPLETION REPORT", title_style))
    
    # Large colored bar
    divider_data = [[""]]
    divider_table = Table(divider_data, colWidths=[500], rowHeights=[4])
    divider_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), primary_color),
        ('BOTTOMPADDING', (0,0), (-1,-1), 0),
        ('TOPPADDING', (0,0), (-1,-1), 0),
    ]))
    story.append(divider_table)
    story.append(Spacer(1, 20))
    
    story.append(Paragraph(report_data.get("program_name", "Untitled Event").upper(), subtitle_style))
    story.append(Spacer(1, 50))
    
    # Metadata Box on Cover Page
    metadata = [
        [Paragraph("Program Name", table_label_style), Paragraph(report_data.get("program_name", "N/A"), body_style)],
        [Paragraph("Program Type", table_label_style), Paragraph(report_data.get("program_type", "N/A"), body_style)],
        [Paragraph("Duration (Days)", table_label_style), Paragraph(str(report_data.get("num_days", "N/A")), body_style)],
        [Paragraph("Target Page Length", table_label_style), Paragraph(f"{report_data.get('target_pages', 'N/A')} Pages", body_style)]
    ]
    meta_table = Table(metadata, colWidths=[180, 320])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), light_bg),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#e2e8f0")),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor("#e2e8f0")),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 8),
        ('BOTTOMPADDING', (0,0), (-1,-1), 8),
        ('LEFTPADDING', (0,0), (-1,-1), 12),
        ('RIGHTPADDING', (0,0), (-1,-1), 12),
    ]))
    story.append(meta_table)
    
    story.append(PageBreak())
    
    # ---------------------------------------------------------
    # TABLE OF CONTENTS
    # ---------------------------------------------------------
    story.append(Paragraph("Table of Contents", h1_style))
    story.append(Spacer(1, 10))
    
    toc_data = [
        [Paragraph("<b>1. Event Overview & Metadata</b>", body_style), Paragraph(". . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . .", body_style), Paragraph("3", body_style)],
        [Paragraph("<b>2. Program Aim & Objectives</b>", body_style), Paragraph(". . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . .", body_style), Paragraph("3", body_style)],
        [Paragraph("<b>3. Program Inclusions</b>", body_style), Paragraph(". . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . .", body_style), Paragraph("3", body_style)],
        [Paragraph("<b>4. Executive Summary (AI Generated)</b>", body_style), Paragraph(". . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . .", body_style), Paragraph("4", body_style)],
        [Paragraph("<b>5. Session Highlights</b>", body_style), Paragraph(". . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . .", body_style), Paragraph("4", body_style)],
        [Paragraph("<b>6. Feedback & Impact Analysis</b>", body_style), Paragraph(". . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . .", body_style), Paragraph("5", body_style)],
        [Paragraph("<b>7. Attached Photos Gallery</b>", body_style), Paragraph(". . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . . .", body_style), Paragraph("6", body_style)],
    ]
    toc_table = Table(toc_data, colWidths=[180, 290, 30])
    toc_table.setStyle(TableStyle([
        ('VALIGN', (0,0), (-1,-1), 'BOTTOM'),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ('TOPPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(toc_table)
    story.append(PageBreak())
    
    # ---------------------------------------------------------
    # SECTION 1: EVENT OVERVIEW
    # ---------------------------------------------------------
    story.append(Paragraph("1. Event Overview & Metadata", h1_style))
    
    overview_details = [
        [Paragraph("<b>Program Name</b>", body_style), Paragraph(report_data.get("program_name", "N/A"), body_style)],
        [Paragraph("<b>Program Type</b>", body_style), Paragraph(report_data.get("program_type", "N/A"), body_style)],
        [Paragraph("<b>Duration (No. of Days)</b>", body_style), Paragraph(str(report_data.get("num_days", "N/A")), body_style)],
        [Paragraph("<b>Target Page Limit</b>", body_style), Paragraph(f"{report_data.get('target_pages', 'N/A')} Pages", body_style)],
    ]
    overview_table = Table(overview_details, colWidths=[180, 320])
    overview_table.setStyle(TableStyle([
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#e2e8f0")),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('TOPPADDING', (0,0), (-1,-1), 6),
        ('BOTTOMPADDING', (0,0), (-1,-1), 6),
        ('LEFTPADDING', (0,0), (-1,-1), 8),
        ('RIGHTPADDING', (0,0), (-1,-1), 8),
    ]))
    story.append(overview_table)
    story.append(Spacer(1, 15))
    
    # ---------------------------------------------------------
    # SECTION 2: AIM & OBJECTIVES
    # ---------------------------------------------------------
    story.append(Paragraph("2. Program Aim & Objectives", h1_style))
    story.append(Paragraph("Aim", h2_style))
    story.append(Paragraph(report_data.get("aim", "No program aim provided."), body_style))
    story.append(Spacer(1, 5))
    story.append(Paragraph("Objectives", h2_style))
    story.append(Paragraph(report_data.get("objective", "No objectives provided."), body_style))
    story.append(Spacer(1, 15))
    
    # ---------------------------------------------------------
    # SECTION 3: PROGRAM INCLUSIONS
    # ---------------------------------------------------------
    story.append(Paragraph("3. Program Inclusions", h1_style))
    story.append(Paragraph(report_data.get("program_inclusions", "No inclusions provided."), body_style))
    story.append(Spacer(1, 15))
    
    # ---------------------------------------------------------
    # SECTION 4: EXECUTIVE SUMMARY (AI GENERATED)
    # ---------------------------------------------------------
    story.append(Paragraph("4. Executive Summary (AI Generated)", h1_style))
    story.append(Paragraph(report_data.get("executive_summary", "No executive summary generated."), body_style))
    story.append(Spacer(1, 15))
    
    # ---------------------------------------------------------
    # SECTION 5: SESSION HIGHLIGHTS
    # ---------------------------------------------------------
    story.append(Paragraph("5. Session Highlights", h1_style))
    story.append(Paragraph(report_data.get("session_highlights", "No session highlights generated."), body_style))
    story.append(Spacer(1, 15))
    
    # ---------------------------------------------------------
    # SECTION 6: FEEDBACK & IMPACT ANALYSIS
    # ---------------------------------------------------------
    story.append(Paragraph("6. Feedback & Impact Analysis", h1_style))
    
    story.append(Paragraph("Student / Participant Feedback Summary", h2_style))
    story.append(Paragraph(report_data.get("student_feedback", "No feedback provided."), body_style))
    story.append(Spacer(1, 10))
    
    story.append(Paragraph("Learning Outcomes", h2_style))
    story.append(Paragraph(report_data.get("learning_outcomes", "No learning outcomes generated."), body_style))
    story.append(Spacer(1, 10))
    
    story.append(Paragraph("Impact Analysis", h2_style))
    story.append(Paragraph(report_data.get("impact_analysis", "No impact analysis generated."), body_style))
    story.append(Spacer(1, 15))
    
    # ---------------------------------------------------------
    # SECTION 7: PHOTO GALLERY
    # ---------------------------------------------------------
    if image_categories:
        story.append(PageBreak())
        story.append(Paragraph("7. Attached Photos Gallery", h1_style))
        
        photo_count = 1
        for category, img_list in image_categories.items():
            if img_list:
                story.append(Paragraph(f"{category} Gallery", h2_style))
                story.append(Spacer(1, 5))
                
                row_images = []
                for idx, img_info in enumerate(img_list):
                    try:
                        rl_img = Image(io.BytesIO(img_info["bytes"]), width=2.2*inch, height=1.65*inch)
                        rl_img.hAlign = 'CENTER'
                        
                        cap_style = ParagraphStyle(
                            f'Cap_{category}_{idx}',
                            parent=body_style,
                            fontSize=8,
                            leading=10,
                            textColor=colors.HexColor("#475569"),
                            alignment=1
                        )
                        caption_para = Paragraph(f"Photo {photo_count}: {img_info['caption']}", cap_style)
                        
                        cell_content = [rl_img, Spacer(1, 3), caption_para, Spacer(1, 10)]
                        row_images.append(cell_content)
                        photo_count += 1
                    except Exception:
                        pass
                
                table_cells = []
                for i in range(0, len(row_images), 2):
                    row = row_images[i:i+2]
                    if len(row) == 1:
                        row.append([])
                    table_cells.append(row)
                
                if table_cells:
                    img_table = Table(table_cells, colWidths=[250, 250])
                    img_table.setStyle(TableStyle([
                        ('VALIGN', (0,0), (-1,-1), 'TOP'),
                        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
                        ('BOTTOMPADDING', (0,0), (-1,-1), 10),
                    ]))
                    story.append(img_table)
                    story.append(Spacer(1, 10))
    
    # Conclusions & Signatures
    story.append(KeepTogether([
        Spacer(1, 15),
        Paragraph("Conclusion", h1_style),
        Paragraph(report_data.get("conclusion", "No conclusion generated."), body_style),
        Spacer(1, 40),
        Table([
            [
                Paragraph("<b>Prepared By:</b><br/><br/><br/>_______________________<br/>Event Coordinator", body_style),
                Paragraph("<b>Approved By:</b><br/><br/><br/>_______________________<br/>Head of Department / Director", body_style)
            ]
        ], colWidths=[250, 250])
    ]))
    
    doc.build(story, canvasmaker=NumberedCanvas)
    buffer.seek(0)
    return buffer.getvalue()


def generate_docx(report_data: dict, logo_bytes: bytes = None, image_categories: dict = None) -> bytes:
    """Generates an editable Word DOCX report."""
    doc = docx.Document()
    
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(10.5)
    style_normal.font.color.rgb = RGBColor(30, 41, 59)
    
    color_primary = RGBColor(30, 58, 138)
    color_secondary = RGBColor(13, 148, 136)
    
    # COVER PAGE
    if logo_bytes:
        try:
            p_logo = doc.add_paragraph()
            p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_logo = p_logo.add_run()
            run_logo.add_picture(io.BytesIO(logo_bytes), width=Inches(1.5))
        except Exception:
            pass
            
    p_title_pre = doc.add_paragraph()
    p_title_pre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_pre = p_title_pre.add_run("EVENT COMPLETION REPORT")
    run_pre.font.size = Pt(24)
    run_pre.font.bold = True
    run_pre.font.color.rgb = color_primary
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run(report_data.get("program_name", "Untitled Event").upper())
    run_title.font.size = Pt(14)
    run_title.font.bold = True
    run_title.font.color.rgb = color_secondary
    
    doc.add_paragraph().paragraph_format.space_before = Pt(40)
    
    # Metadata Table
    table = doc.add_table(rows=4, cols=2)
    table.alignment = docx.enum.table.WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    meta_rows = [
        ("Program Name:", report_data.get("program_name", "N/A")),
        ("Program Type:", report_data.get("program_type", "N/A")),
        ("Duration (No. of Days):", str(report_data.get("num_days", "N/A"))),
        ("Target Page Limit:", f"{report_data.get('target_pages', 'N/A')} Pages"),
    ]
    
    for idx, (label, val) in enumerate(meta_rows):
        row = table.rows[idx]
        row.cells[0].paragraphs[0].add_run(label).bold = True
        row.cells[0].paragraphs[0].runs[0].font.color.rgb = color_primary
        row.cells[1].paragraphs[0].add_run(val)
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(4.3)
        
    doc.add_page_break()
    
    def add_section_heading(text, level=1):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(text)
        run.bold = True
        if level == 1:
            run.font.size = Pt(16)
            run.font.color.rgb = color_primary
        else:
            run.font.size = Pt(12)
            run.font.color.rgb = color_secondary
        return p
        
    # Table of Contents placeholder
    add_section_heading("Table of Contents", 1)
    doc.add_paragraph("1. Event Overview & Metadata\n2. Program Aim & Objectives\n3. Program Inclusions\n4. Executive Summary (AI Generated)\n5. Session Highlights\n6. Feedback & Impact Analysis\n7. Photo Gallery\n8. Conclusion & Signatures")
    doc.add_page_break()
    
    # SECTION 1
    add_section_heading("1. Event Overview & Metadata", 1)
    tbl_details = doc.add_table(rows=4, cols=2)
    tbl_details.alignment = docx.enum.table.WD_TABLE_ALIGNMENT.CENTER
    tbl_details.autofit = False
    
    details_rows = [
        ("Program Name", report_data.get("program_name", "N/A")),
        ("Program Type", report_data.get("program_type", "N/A")),
        ("Duration (Days)", str(report_data.get("num_days", "N/A"))),
        ("Target Page Limit", f"{report_data.get('target_pages', 'N/A')} Pages"),
    ]
    
    for idx, (label, val) in enumerate(details_rows):
        row = tbl_details.rows[idx]
        row.cells[0].paragraphs[0].add_run(label).bold = True
        row.cells[1].paragraphs[0].add_run(val)
        row.cells[0].width = Inches(2.2)
        row.cells[1].width = Inches(4.3)
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(r'<w:tcBorders %s><w:top w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/><w:left w:val="none"/><w:bottom w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/><w:right w:val="none"/></w:tcBorders>' % nsdecls('w'))
            tcPr.append(tcBorders)
            
    doc.add_paragraph().paragraph_format.space_before = Pt(10)
    
    # SECTION 2
    add_section_heading("2. Program Aim & Objectives", 1)
    add_section_heading("Aim", 2)
    doc.add_paragraph(report_data.get("aim", ""))
    add_section_heading("Objectives", 2)
    doc.add_paragraph(report_data.get("objective", ""))
    
    # SECTION 3
    add_section_heading("3. Program Inclusions", 1)
    doc.add_paragraph(report_data.get("program_inclusions", ""))
    
    # SECTION 4
    add_section_heading("4. Executive Summary (AI Generated)", 1)
    doc.add_paragraph(report_data.get("executive_summary", ""))
    
    # SECTION 5
    add_section_heading("5. Session Highlights", 1)
    doc.add_paragraph(report_data.get("session_highlights", ""))
    
    # SECTION 6
    add_section_heading("6. Feedback & Impact Analysis", 1)
    add_section_heading("Student Feedback", 2)
    doc.add_paragraph(report_data.get("student_feedback", ""))
    add_section_heading("Learning Outcomes", 2)
    doc.add_paragraph(report_data.get("learning_outcomes", ""))
    add_section_heading("Impact Analysis", 2)
    doc.add_paragraph(report_data.get("impact_analysis", ""))
    
    # SECTION 7
    if image_categories:
        doc.add_page_break()
        add_section_heading("7. Event Photo Gallery", 1)
        photo_idx = 1
        for category, img_list in image_categories.items():
            if img_list:
                add_section_heading(f"{category} Photos", 2)
                for item in img_list:
                    try:
                        p_img = doc.add_paragraph()
                        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_img.paragraph_format.keep_with_next = True
                        run_img = p_img.add_run()
                        run_img.add_picture(io.BytesIO(item["bytes"]), width=Inches(3.5))
                        
                        p_cap = doc.add_paragraph()
                        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_cap.paragraph_format.space_after = Pt(12)
                        run_cap = p_cap.add_run(f"Photo {photo_idx}: {item['caption']}")
                        run_cap.italic = True
                        run_cap.font.size = Pt(9)
                        
                        photo_idx += 1
                    except Exception:
                        pass
                        
    # CONCLUSION
    add_section_heading("Conclusion", 1)
    doc.add_paragraph(report_data.get("conclusion", ""))
    
    doc.add_paragraph().paragraph_format.space_before = Pt(30)
    
    sig_table = doc.add_table(rows=1, cols=2)
    sig_table.alignment = docx.enum.table.WD_TABLE_ALIGNMENT.CENTER
    sig_table.autofit = False
    
    sig_table.rows[0].cells[0].paragraphs[0].add_run("Prepared By:\n\n\n_______________________\nEvent Coordinator")
    sig_table.rows[0].cells[1].paragraphs[0].add_run("Approved By:\n\n\n_______________________\nHead of Department / Director")
    sig_table.rows[0].cells[0].width = Inches(3.2)
    sig_table.rows[0].cells[1].width = Inches(3.2)
    
    doc_buffer = io.BytesIO()
    doc.save(doc_buffer)
    doc_buffer.seek(0)
    return doc_buffer.getvalue()

# ---------------------------------------------------------
# GEMINI ENGINE MODULE (Inline ai_processor.py)
# ---------------------------------------------------------

class AIProcessor:
    def __init__(self, api_key=None):
        self.api_key = api_key or os.getenv("GEMINI_API_KEY")
        if self.api_key:
            genai.configure(api_key=self.api_key)

    def configure_key(self, api_key):
        self.api_key = api_key
        genai.configure(api_key=api_key)

    def generate_report_sections(self, metadata: dict, form_inputs: dict, documents_text: str) -> dict:
        if not self.api_key:
            raise ValueError("Gemini API key is not configured.")

        event_details = f"""
Program Name: {metadata.get('program_name', 'N/A')}
Program Type: {metadata.get('program_type', 'N/A')}
No of Days: {metadata.get('num_days', 'N/A')}
Target Page Limit: {metadata.get('target_pages', 'N/A')} Pages

Manually Inputted Details:
Program Aim: {form_inputs.get('aim', 'N/A')}
Program Objective: {form_inputs.get('objective', 'N/A')}
Program Inclusions: {form_inputs.get('program_inclusions', 'N/A')}

Uploaded Documents Extracted Content:
{documents_text[:12000]}
"""

        system_instruction = (
            "You are a professional report writer for educational and corporate events. "
            "Based on the provided event details, objectives, aim, inclusions, and extracted documents content, you must generate "
            "5 specific sections of the final report. "
            "Keep the language extremely professional, polished, coherent, and formal. "
            "Write the sections in clear, detailed paragraphs. Do not use markdown headers in your response; "
            "instead, format the output using the exact XML tags specified below for each section."
        )

        user_prompt = f"""
Synthesize the event details and documents, and write the following sections for the report. Wrap each section inside the corresponding XML tags:

1. <executive_summary>
Generate a comprehensive, formal, and professional summary of the event (1-2 paragraphs), highlighting the program's purpose (Aim & Objective), background, execution, and general success.
</executive_summary>

2. <session_highlights>
Summarize the key sessions, discussions, and topics presented during the event. Synthesize this based on the uploaded documents and notes (2 paragraphs).
</session_highlights>

3. <learning_outcomes>
Generate a list of clear, measurable learning outcomes that the participants achieved. Express them in a clean narrative format (1-2 paragraphs).
</learning_outcomes>

4. <impact_analysis>
Describe how the participants/students benefited from this event. Address their skill improvement, domain knowledge expansion, and feedback response (1-2 paragraphs).
</impact_analysis>

5. <conclusion>
Generate a strong, formal concluding paragraph summarizing the key impact and expressing gratitude or future outlook (1 paragraph).
</conclusion>

Here are all the details to base your report on:
{event_details}
"""

        model = genai.GenerativeModel("gemini-1.5-flash", system_instruction=system_instruction)
        response = model.generate_content(user_prompt)
        text = response.text

        sections = {}
        for tag in ["executive_summary", "session_highlights", "learning_outcomes", "impact_analysis", "conclusion"]:
            start_tag = f"<{tag}>"
            end_tag = f"</{tag}>"
            if start_tag in text and end_tag in text:
                sections[tag] = text.split(start_tag)[1].split(end_tag)[0].strip()
            else:
                sections[tag] = f"Unable to automatically generate {tag.replace('_', ' ')}. Please edit this field manually."
        return sections

    def generate_image_caption(self, image_bytes: bytes) -> str:
        if not self.api_key:
            return "Event photo showing activities."

        try:
            img = PIL.Image.open(io.BytesIO(image_bytes))
            if img.width > 1200 or img.height > 1200:
                img.thumbnail((1200, 1200))
                
            model = genai.GenerativeModel("gemini-1.5-flash")
            
            prompt = (
                "Provide a professional, formal caption for this event photo to be used in an official report. "
                "The caption must be concise (1 sentence). Do not start with filler words like 'This image shows' or 'A photo of'. "
                "Focus on the professional activity happening in the photo."
            )
            
            response = model.generate_content([prompt, img])
            return response.text.strip()
        except Exception:
            return "Event session photo."

# ---------------------------------------------------------
# STREAMLIT INTERFACE CONFIGURATION
# ---------------------------------------------------------

# Custom premium CSS styling (Premium Black and Gold/Yellow Theme)
st.markdown("""
<style>
    @import url('https://fonts.googleapis.com/css2?family=Outfit:wght@300;400;600;700&display=swap');
    
    html, body, [data-testid="stAppViewContainer"] {
        font-family: 'Outfit', sans-serif;
        background-color: #050505;
        color: #f3f4f6;
    }
    
    /* Header Container Gradient */
    .header-container {
        background: linear-gradient(135deg, #151515 0%, #000000 100%);
        padding: 2.5rem;
        border-radius: 16px;
        border: 1px solid rgba(255, 204, 0, 0.25);
        box-shadow: 0 8px 32px 0 rgba(255, 204, 0, 0.05);
        margin-bottom: 2rem;
    }
    .header-title {
        font-size: 2.6rem;
        font-weight: 700;
        background: linear-gradient(to right, #ffe600, #ffb300, #ff8800);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        margin-bottom: 0.5rem;
    }
    .header-subtitle {
        color: #d1d5db;
        font-size: 1.1rem;
    }
    
    /* Sidebar Styling */
    [data-testid="stSidebar"] {
        background-color: #0b0b0b !important;
        border-right: 1px solid rgba(255, 204, 0, 0.15);
    }
    
    /* Card Styles */
    .panel-card {
        background: rgba(20, 20, 20, 0.7);
        backdrop-filter: blur(12px);
        border: 1px solid rgba(255, 204, 0, 0.15);
        border-radius: 12px;
        padding: 1.5rem;
        margin-bottom: 1rem;
    }
    
    .status-active {
        color: #ffcc00;
        font-weight: 600;
    }
    .status-inactive {
        color: #f43f5e;
        font-weight: 600;
    }
    
    /* Custom buttons */
    div.stButton > button {
        background: linear-gradient(to right, #ffcc00, #ff9900);
        color: #000000 !important;
        border: none;
        padding: 0.6rem 1.8rem;
        border-radius: 8px;
        font-weight: 700;
        transition: all 0.3s ease;
    }
    div.stButton > button:hover {
        transform: translateY(-2px);
        box-shadow: 0 4px 15px rgba(255, 204, 0, 0.3);
        background: linear-gradient(to right, #ffe066, #ffaa00);
        color: #000000 !important;
    }
    
    /* Preview Section Custom Styles */
    .preview-doc {
        background-color: #ffffff;
        color: #1a1a1a;
        padding: 40px;
        border-radius: 8px;
        box-shadow: 0 10px 25px rgba(0,0,0,0.1);
        border: 1px solid #e2e8f0;
        max-width: 800px;
        margin: 0 auto;
        font-family: Arial, sans-serif;
    }
    .preview-header {
        border-bottom: 2px solid #ffcc00;
        padding-bottom: 10px;
        margin-bottom: 20px;
    }
    .preview-title {
        color: #000000;
        font-size: 24px;
        font-weight: bold;
    }
    .preview-section-title {
        color: #ff9900;
        font-size: 18px;
        font-weight: bold;
        border-bottom: 1px solid #e2e8f0;
        margin-top: 25px;
        padding-bottom: 5px;
    }
    .preview-sub-title {
        color: #cc9900;
        font-size: 14px;
        font-weight: bold;
        margin-top: 15px;
    }
    .preview-table {
        width: 100%;
        border-collapse: collapse;
        margin: 15px 0;
    }
    .preview-table th, .preview-table td {
        border: 1px solid #cbd5e1;
        padding: 8px;
        text-align: left;
        font-size: 13px;
    }
    .preview-table th {
        background-color: #f8fafc;
        color: #000000;
    }
    .preview-img-container {
        display: inline-block;
        width: 45%;
        margin: 2%;
        text-align: center;
        vertical-align: top;
        border: 1px solid #cbd5e1;
        padding: 5px;
        border-radius: 4px;
        background: #f8fafc;
    }
    .preview-img-caption {
        font-size: 11px;
        color: #64748b;
        font-style: italic;
        margin-top: 4px;
    }
</style>
""", unsafe_allow_html=True)

# Session state initialization
if "event_data" not in st.session_state:
    st.session_state.event_data = {
        "program_name": "National Seminar on Generative Artificial Intelligence",
        "num_days": 3,
        "program_type": "Seminar",
        "aim": "To introduce academic researchers and industry professionals to generative AI prompt engineering, large language model design, and multi-agent systems.",
        "objective": "1. Provide insights into Gemini model API integration.\n2. Conduct practical labs for document indexing and custom search systems.\n3. Discuss security controls in cloud-based API architectures.",
        "program_inclusions": "1. Access to practical lab notebooks.\n2. Course participation certificate.\n3. Institutional lunch, tea break refreshments, and resource toolkit.",
        "student_feedback": "Highly positive feedback. Participants valued the hands-on lab sessions and clarity of topic delivery.",
        "target_pages": 4,
        "executive_summary": "",
        "session_highlights": "",
        "learning_outcomes": "",
        "impact_analysis": "",
        "conclusion": "",
        "documents_text": ""
    }

if "logo_bytes" not in st.session_state:
    st.session_state.logo_bytes = None

if "image_categories" not in st.session_state:
    st.session_state.image_categories = {
        "Attached Event Photos": []
    }

if "api_key" not in st.session_state:
    st.session_state.api_key = os.getenv("GEMINI_API_KEY") or ""

# Sidebar Configuration
with st.sidebar:
    st.image("https://img.icons8.com/nolan/96/news.png", width=60)
    st.markdown("### 🤖 Report Engine")
    st.write("Automatically draft executive-level event summaries with AI & professional exports.")
    
    st.markdown("---")
    
    menu = st.radio(
        "Navigation Menu",
        ["🏠 Dashboard", "📝 Event Inputs", "📷 Upload Photos", "✨ AI Refiner & Editor", "🔍 Preview & Export"]
    )
    
    st.markdown("---")
    
    st.markdown("#### 🔑 Gemini API Settings")
    if st.session_state.api_key:
        st.markdown('Status: <span class="status-active">🟢 Active</span>', unsafe_allow_html=True)
        if st.button("Change API Key"):
            st.session_state.api_key = ""
            st.rerun()
    else:
        st.markdown('Status: <span class="status-inactive">🔴 Missing</span>', unsafe_allow_html=True)
        key_input = st.text_input("Enter Gemini API Key", type="password")
        if key_input:
            st.session_state.api_key = key_input
            st.success("Key updated!")
            st.rerun()

processor = AIProcessor(api_key=st.session_state.api_key)

# Dashboard
if menu == "🏠 Dashboard":
    st.markdown("""
    <div class="header-container">
        <div class="header-title">AI Event Report Generator</div>
        <div class="header-subtitle">Streamlined completion reporting with automated document summarization, photo captioning, and exports.</div>
    </div>
    """, unsafe_allow_html=True)
    
    col1, col2 = st.columns(2)
    with col1:
        st.markdown('<div class="panel-card">', unsafe_allow_html=True)
        st.markdown("### 🚀 Quick Start Guide")
        st.markdown("""
        Follow these simple steps to build your official report:
        1. **Fill Event Inputs**: Provide program name, number of days, type of program, and requirements (Aim, Objective, Inclusions, Target Pages).
        2. **Upload Documents**: Attach files regarding the program to add context for AI summaries.
        3. **Upload Photos**: Upload available photos to be attached to the final report.
        4. **Generate with AI**: Let Gemini read inputs and documents to write structured summaries and captions.
        5. **Refine & Export**: Preview the report and download PDF or Word exports.
        """)
        st.markdown('</div>', unsafe_allow_html=True)
        
    with col2:
        st.markdown('<div class="panel-card">', unsafe_allow_html=True)
        st.markdown("### 📊 Active Report Details")
        st.write(f"**Program Name:** {st.session_state.event_data['program_name']}")
        st.write(f"**Program Type:** {st.session_state.event_data['program_type']}")
        st.write(f"**Duration (Days):** {st.session_state.event_data['num_days']}")
        
        total_photos = len(st.session_state.image_categories["Attached Event Photos"])
        st.write(f"**Uploaded Photos:** {total_photos} items")
        
        has_ai = st.session_state.event_data['executive_summary'] != ""
        st.write(f"**AI Sections Synthesized:** {'🟢 Yes' if has_ai else '🔴 No'}")
        st.markdown('</div>', unsafe_allow_html=True)

# Event Inputs
elif menu == "📝 Event Inputs":
    st.markdown("## 📝 Program Information & Manual Entries")
    st.write("Provide details and upload files relating to the program.")
    
    tab1, tab2, tab3 = st.tabs(["🏛️ Metadata & Requirements", "📂 Documents Uploader", "✍️ Manual Contents"])
    
    with tab1:
        st.markdown('<div class="panel-card">', unsafe_allow_html=True)
        col1, col2 = st.columns(2)
        with col1:
            st.session_state.event_data["program_name"] = st.text_input("Program Name", st.session_state.event_data["program_name"])
            st.session_state.event_data["program_type"] = st.selectbox(
                "Type of Program", 
                ["Seminar", "Workshop", "Conference", "FDP", "Webinar", "Symposium", "Guest Lecture", "Field Visit"], 
                index=0
            )
        with col2:
            st.session_state.event_data["num_days"] = st.number_input("No. of Days", min_value=1, value=int(st.session_state.event_data["num_days"]))
            st.session_state.event_data["target_pages"] = st.number_input("No. of Pages (Target Length)", min_value=1, value=int(st.session_state.event_data["target_pages"]))
        st.markdown('</div>', unsafe_allow_html=True)
        
    with tab2:
        st.markdown('<div class="panel-card">', unsafe_allow_html=True)
        col1, col2 = st.columns(2)
        with col1:
            logo_file = st.file_uploader("Upload Institution Logo (PNG/JPG)", type=["png", "jpg", "jpeg"])
            if logo_file:
                st.session_state.logo_bytes = logo_file.read()
                st.image(st.session_state.logo_bytes, width=120, caption="Uploaded Logo Preview")
                
        with col2:
            uploaded_docs = st.file_uploader("Upload Documents regarding the Program (PDF/DOCX) (Multiple)", type=["pdf", "docx"], accept_multiple_files=True)
            
            if uploaded_docs and st.button("Parse Uploaded Documents"):
                with st.spinner("Extracting contents..."):
                    texts = []
                    for f in uploaded_docs:
                        name = f.name
                        bytes_data = f.read()
                        if name.endswith(".pdf"):
                            texts.append(f"--- DOCUMENT: {name} ---\n" + extract_text_from_pdf(bytes_data))
                        elif name.endswith(".docx"):
                            texts.append(f"--- DOCUMENT: {name} ---\n" + extract_text_from_docx(bytes_data))
                            
                    combined_docs = "\n\n".join(texts)
                    if combined_docs:
                        st.session_state.event_data["documents_text"] = combined_docs
                        st.success("Successfully parsed text from documents! Content added to AI context.")
                    else:
                        st.info("No text files were parsed.")
        st.markdown('</div>', unsafe_allow_html=True)
        
    with tab3:
        st.markdown('<div class="panel-card">', unsafe_allow_html=True)
        st.session_state.event_data["aim"] = st.text_area("Aim of the Program", st.session_state.event_data["aim"], height=100)
        st.session_state.event_data["objective"] = st.text_area("Objectives of the Program", st.session_state.event_data["objective"], height=120)
        st.session_state.event_data["program_inclusions"] = st.text_area("Program Inclusions (e.g. materials, certificates, meals)", st.session_state.event_data["program_inclusions"], height=100)
        st.session_state.event_data["student_feedback"] = st.text_area("Student / Participant Feedback Summary", st.session_state.event_data.get("student_feedback", ""), height=100)
        st.markdown('</div>', unsafe_allow_html=True)

# Upload Photos
elif menu == "📷 Upload Photos":
    st.markdown("## 📷 Event Photos Upload & Gallery")
    st.write("Upload available photos to attach regarding the program.")
    
    st.markdown('<div class="panel-card">', unsafe_allow_html=True)
    uploaded_photos = st.file_uploader(
        "Upload Program Photos (Multiple)", 
        type=["png", "jpg", "jpeg"], 
        accept_multiple_files=True
    )
    if uploaded_photos:
        for f in uploaded_photos:
            existing_names = [img["name"] for img in st.session_state.image_categories["Attached Event Photos"]]
            if f.name not in existing_names:
                img_bytes = f.read()
                st.session_state.image_categories["Attached Event Photos"].append({
                    "name": f.name,
                    "bytes": img_bytes,
                    "caption": "Program activities photo."
                })
        st.success(f"Added {len(uploaded_photos)} photos to gallery!")
        
    st.markdown("---")
    
    current_list = st.session_state.image_categories["Attached Event Photos"]
    if current_list:
        cols = st.columns(3)
        for idx, item in enumerate(current_list):
            col_idx = idx % 3
            with cols[col_idx]:
                st.image(item["bytes"], use_container_width=True)
                new_cap = st.text_input(
                    f"Caption #{idx+1}", 
                    value=item["caption"], 
                    key=f"cap_edit_photo_{idx}"
                )
                item["caption"] = new_cap
                if st.button("🗑️ Delete", key=f"del_photo_{idx}"):
                    st.session_state.image_categories["Attached Event Photos"].pop(idx)
                    st.rerun()
    else:
        st.info("No photos uploaded yet.")
    st.markdown('</div>', unsafe_allow_html=True)

# AI Refiner & Editor
elif menu == "✨ AI Refiner & Editor":
    st.markdown("## ✨ AI Automated Content Refinement")
    st.write("Generate structured summaries and descriptions using Gemini, or use local template heuristics, then edit as required.")
    
    if not st.session_state.api_key:
        st.info("💡 **Tip**: Enter a Gemini API Key in the sidebar to use advanced AI generation and automatic photo captioning.")
    
    col1, col2 = st.columns(2)
    
    with col1:
        st.markdown("### 📝 Text Report Sections")
        if st.session_state.api_key:
            if st.button("⚡ Generate AI Report (Gemini API)"):
                with st.spinner("Generating report sections..."):
                    try:
                        results = processor.generate_report_sections(
                            st.session_state.event_data,
                            st.session_state.event_data,
                            st.session_state.event_data.get("documents_text", "")
                        )
                        for key in ["executive_summary", "session_highlights", "learning_outcomes", "impact_analysis", "conclusion"]:
                            st.session_state.event_data[key] = results.get(key, "")
                        st.success("Successfully generated report text blocks! Review them below.")
                    except Exception as e:
                        st.error(f"Failed to generate reports: {e}")
        
        if st.button("📝 Generate Smart Templates (Local Heuristics)"):
            data = st.session_state.event_data
            aim = data.get("aim", "")
            objective = data.get("objective", "")
            inclusions = data.get("program_inclusions", "")
            name = data.get("program_name", "the program")
            type_p = data.get("program_type", "event")
            days = data.get("num_days", 1)
            
            data["executive_summary"] = (
                f"The primary aim of this {type_p} was to address the following: '{aim}'. Over a span of {days} days, the program was "
                f"successfully organized and completed. Key objectives included achieving: \n{objective}\n\nParticipants reported a high "
                f"degree of satisfaction with the learning pace and overall session structure."
            )
            data["session_highlights"] = (
                f"During this {days}-day {type_p}, speakers and coordinators focused on key topics related to: '{aim}'. "
                f"Interactive sessions were conducted to help participants understand practical execution. Notable inclusions "
                f"like '{inclusions}' ensured that attendees had all the resources required for a productive environment."
            )
            data["learning_outcomes"] = (
                f"By the conclusion of the {type_p}, participants achieved the following target competencies:\n"
                f"1. Enhanced understanding of core principles linked to: '{aim}'.\n"
                f"2. Ability to implement practical methodologies to accomplish: \n{objective}\n"
                f"3. Collaborative networking and resource sharing among peers."
            )
            data["impact_analysis"] = (
                f"The program created a strong foundation for future applications. Based on the aim, the participants "
                f"now possess the skills required to address relevant domain challenges. The structured inclusions and "
                f"comprehensive objectives set a benchmark for future iterations of similar educational workshops."
            )
            data["conclusion"] = (
                f"In conclusion, the {days}-day {type_p} on '{name}' successfully met all its defined objectives. We extend our sincere "
                f"gratitude to the resource persons, organizational coordinators, and participants whose dedication made this event a resounding success."
            )
            st.success("Successfully filled sections with local smart templates! Review them below.")
            st.rerun()

    with col2:
        st.markdown("### 📷 Photo Captions")
        if st.session_state.api_key:
            if st.button("📷 Generate AI Captions"):
                with st.spinner("AI describing photos..."):
                    processed_count = 0
                    for img in st.session_state.image_categories["Attached Event Photos"]:
                        try:
                            caption = processor.generate_image_caption(img["bytes"])
                            img["caption"] = caption
                            processed_count += 1
                        except Exception:
                            pass
                    if processed_count > 0:
                        st.success(f"Generated captions for {processed_count} photos!")
                        st.rerun()
                    else:
                        st.info("No photos found or API error occurred.")
        
        if st.button("📷 Set Default Photo Captions"):
            processed_count = 0
            for idx, img in enumerate(st.session_state.image_categories["Attached Event Photos"]):
                img["caption"] = f"Photograph showing activities during session {idx+1} of the {st.session_state.event_data.get('program_name')}."
                processed_count += 1
            if processed_count > 0:
                st.success(f"Updated default captions for {processed_count} photos!")
                st.rerun()
            else:
                st.info("No photos found.")

    st.markdown("---")
    
    st.subheader("Edit Report Content")
    st.session_state.event_data["executive_summary"] = st.text_area("Executive Summary", st.session_state.event_data["executive_summary"], height=120)
    st.session_state.event_data["session_highlights"] = st.text_area("Session Highlights", st.session_state.event_data["session_highlights"], height=120)
    st.session_state.event_data["learning_outcomes"] = st.text_area("Learning Outcomes", st.session_state.event_data["learning_outcomes"], height=120)
    st.session_state.event_data["impact_analysis"] = st.text_area("Impact Analysis", st.session_state.event_data["impact_analysis"], height=120)
    st.session_state.event_data["conclusion"] = st.text_area("Conclusion", st.session_state.event_data["conclusion"], height=120)

# Preview and Export
elif menu == "🔍 Preview & Export":
    st.markdown("## 🔍 Report Preview & Document Generation")
    st.write("Preview the layout and download PDF/Word exports.")
    
    col1, col2 = st.columns(2)
    with col1:
        pdf_bytes = generate_pdf(st.session_state.event_data, st.session_state.logo_bytes, st.session_state.image_categories)
        st.download_button(
            label="📥 Download PDF Report",
            data=pdf_bytes,
            file_name="event_report.pdf",
            mime="application/pdf",
            use_container_width=True
        )
        
    with col2:
        docx_bytes = generate_docx(st.session_state.event_data, st.session_state.logo_bytes, st.session_state.image_categories)
        st.download_button(
            label="📥 Download DOCX Report",
            data=docx_bytes,
            file_name="event_report.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            use_container_width=True
        )
        
    st.markdown("---")
    st.subheader("Interactive HTML Preview")
    
    st.markdown(f"""
    <div class="preview-doc">
        <div class="preview-header" style="text-align: center;">
            {"<div style='font-size: 11px; font-weight: bold; color: #1e3a8a;'>INSTITUTION LOGO PLACEHOLDER</div>" if not st.session_state.logo_bytes else ""}
            <div class="preview-title">EVENT COMPLETION REPORT</div>
            <div style="font-size: 16px; font-weight: bold; color: #0d9488; text-transform: uppercase; margin-top: 10px;">
                {st.session_state.event_data['program_name']}
            </div>
        </div>
        
        <table class="preview-table">
            <tr><th width="40%">Field</th><th width="60%">Value</th></tr>
            <tr><td><b>Program Name</b></td><td>{st.session_state.event_data['program_name']}</td></tr>
            <tr><td><b>Program Type</b></td><td>{st.session_state.event_data['program_type']}</td></tr>
            <tr><td><b>Duration (No. of Days)</b></td><td>{st.session_state.event_data['num_days']} Days</td></tr>
            <tr><td><b>Target Page Limit</b></td><td>{st.session_state.event_data['target_pages']} Pages</td></tr>
        </table>
        
        <div class="preview-section-title">1. Aim & Objectives</div>
        <p style="font-size: 13px; line-height: 1.5; color: #475569;">
            <b>Aim:</b><br/>{st.session_state.event_data['aim']}<br/><br/>
            <b>Objectives:</b><br/>{st.session_state.event_data['objective']}
        </p>
        
        <div class="preview-section-title">2. Program Inclusions</div>
        <p style="font-size: 13px; line-height: 1.5; color: #475569; white-space: pre-wrap;">
            {st.session_state.event_data['program_inclusions']}
        </p>
        
        <div class="preview-section-title">3. Executive Summary (AI Generated)</div>
        <p style="font-size: 13px; line-height: 1.5; color: #475569;">
            {st.session_state.event_data['executive_summary'] or "<i>No executive summary generated yet. Navigate to 'AI Refiner' to generate.</i>"}
        </p>
        
        <div class="preview-section-title">4. Session Highlights</div>
        <p style="font-size: 13px; line-height: 1.5; color: #475569;">
            {st.session_state.event_data['session_highlights'] or "<i>No session highlights generated.</i>"}
        </p>
        
        <div class="preview-section-title">5. Feedback & Outcomes</div>
        <p style="font-size: 13px; line-height: 1.5; color: #475569;">
            <b>Feedback Summary:</b><br/>{st.session_state.event_data['student_feedback']}<br/><br/>
            <b>Learning Outcomes:</b><br/>{st.session_state.event_data['learning_outcomes'] or "<i>No learning outcomes generated yet.</i>"}<br/><br/>
            <b>Impact Analysis:</b><br/>{st.session_state.event_data['impact_analysis'] or "<i>No impact analysis generated yet.</i>"}
        </p>
        
        <div class="preview-section-title">6. Conclusion</div>
        <p style="font-size: 13px; line-height: 1.5; color: #475569;">
            {st.session_state.event_data['conclusion'] or "<i>No conclusion generated yet.</i>"}
        </p>
        
        <div class="preview-section-title">7. Signatures</div>
        <table style="width: 100%; margin-top: 30px; border: none;">
            <tr style="border: none;">
                <td style="border: none; font-size: 13px;">
                    <b>Prepared By:</b><br/><br/><br/>_______________________<br/>Event Coordinator
                </td>
                <td style="border: none; font-size: 13px; text-align: right;">
                    <b>Approved By:</b><br/><br/><br/>_______________________<br/>Head of Department / Director
                </td>
            </tr>
        </table>
    </div>
    """, unsafe_allow_html=True)
    
    current_photos = st.session_state.image_categories["Attached Event Photos"]
    if current_photos:
        st.markdown("<div style='max-width: 800px; margin: 20px auto;'><div class='preview-section-title'>Photo Gallery</div></div>", unsafe_allow_html=True)
        cols = st.columns(3)
        for idx, img_data in enumerate(current_photos):
            col_i = idx % 3
            with cols[col_i]:
                st.image(img_data["bytes"], caption=f"Photo {idx+1}: {img_data['caption']}", use_container_width=True)
